import argparse
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(r"C:\TaxCaseManager")
TEMPLATE_DOCX_DIR = BASE_DIR / "templates" / "client_forms_docx"
OUTPUT_DOCX_DIR = BASE_DIR / "outputs" / "client_forms_docx"
CASES_DIR = BASE_DIR / "cases"
LOGS_DIR = BASE_DIR / "logs"

NOTICE_TEXT = "본 서류는 양도세 사전진단을 위한 기초자료이며, 최종 세무 판단은 세무사 검토가 필요합니다."


FORM_FILES = {
    "01_client_tax_consult_application.docx": "양도세 사전진단 상담 신청서",
    "02_property_basic_info.docx": "양도 부동산 기본정보 확인서",
    "03_household_house_status.docx": "주택 보유 및 세대원 확인서",
    "04_residence_period_check.docx": "거주기간 확인서",
    "05_acquisition_price_expense_check.docx": "취득가액 및 필요경비 증빙 확인서",
    "06_exception_issue_checklist.docx": "양도세 예외사항 체크리스트",
    "07_privacy_and_tax_accountant_consent.docx": "개인정보 수집 및 세무사 검토 전달 동의서",
    "08_required_documents_checklist.docx": "양도세 사전진단 필요서류 체크리스트",
    "09_final_client_confirmation.docx": "고객 최종 확인서",
    "10_client_truthful_disclosure_notice.docx": "고객 제공자료 성실확인 및 불이익 고지서",
    "11_tax_filing_delegation_and_identity_check.docx": "양도소득세 신고대행 위임 및 본인확인 안내서",
    "12_tax_accountant_fee_scope_agreement.docx": "세무사 신고대행 수수료 및 업무범위 확인서",
    "13_final_tax_filing_confirmation.docx": "최종 신고내용 확인서",
}


def write_log(message):
    LOGS_DIR.mkdir(exist_ok=True)
    log_path = LOGS_DIR / "client_forms_docx_log.txt"
    now = datetime.now().isoformat(timespec="seconds")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{now}] {message}\n")


def set_cell_shading(cell, fill="EDEDED"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_doc_style(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)


def add_title(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(16)
    doc.add_paragraph("")


def add_notice(doc):
    p = doc.add_paragraph()
    run = p.add_run("※ " + NOTICE_TEXT)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(9)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(12)


def add_field_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "항목", bold=True)
    set_cell_text(hdr[1], "작성 내용", bold=True)
    set_cell_shading(hdr[0])
    set_cell_shading(hdr[1])

    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], value)

    doc.add_paragraph("")
    return table


def add_check_items(doc, items):
    for item in items:
        p = doc.add_paragraph()
        run = p.add_run(f"□ {item}")
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        run.font.size = Pt(10.5)
    doc.add_paragraph("")


def add_signature_block(doc):
    doc.add_paragraph("")
    add_field_table(doc, [
        ("고객명", "____________________________"),
        ("서명", "____________________________"),
        ("작성일", "________년 ________월 ________일"),
    ])


def add_common_header_table(doc):
    add_field_table(doc, [
        ("사건번호", ""),
        ("고객명", ""),
        ("연락처", ""),
        ("부동산 사무실", ""),
        ("담당자", ""),
    ])


def build_01(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 상담 목적")
    add_check_items(doc, [
        "매도 전 세금 확인",
        "계약 전 검토",
        "계약 후 잔금 전 검토",
        "세무사 검토 요청",
        "신고대행 연결 희망",
        "기타",
    ])
    add_section_heading(doc, "2. 현재 진행상태")
    add_check_items(doc, [
        "매도 검토 중",
        "매매계약 체결 전",
        "계약 체결 후 잔금 전",
        "이미 양도 완료",
        "기타",
    ])
    add_section_heading(doc, "3. 확인사항")
    add_check_items(doc, [
        "본 리포트는 세무 판단 확정이 아님을 확인했습니다.",
        "모르는 항목은 임의로 작성하지 않고 모름/확인 필요로 표시하겠습니다.",
        "세무사 검토가 필요한 경우 별도 비용이 발생할 수 있음을 확인했습니다.",
    ])
    add_signature_block(doc)


def build_02(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 매도 대상 부동산")
    add_field_table(doc, [
        ("부동산 주소", ""),
        ("부동산 종류", "□ 아파트 □ 단독주택 □ 다세대/빌라 □ 오피스텔 □ 토지 □ 농지 □ 임야 □ 상가 □ 겸용주택 □ 기타"),
        ("전용면적/토지면적", ""),
        ("건물면적", ""),
        ("소유 형태", "□ 단독소유 □ 부부 공동명의 □ 가족 공동명의 □ 일부지분 □ 기타"),
        ("본인 지분율", ""),
        ("공동명의자 및 지분율", ""),
        ("취득 경위", "□ 매매 □ 상속 □ 증여 □ 분양 □ 재개발/재건축 □ 교환 □ 기타"),
        ("취득일", ""),
        ("등기접수일", ""),
        ("실제 잔금일 또는 실제 취득일", ""),
        ("양도 예정일", ""),
        ("잔금 예정일", ""),
        ("양도 예정가액 또는 계약금액", ""),
        ("양도 상태", "□ 예정 □ 계약 전 □ 계약 후 잔금 전 □ 양도 완료"),
    ])
    add_section_heading(doc, "2. 추가 확인사항")
    add_check_items(doc, [
        "취득일과 등기일이 다를 수 있음",
        "잔금일이 정확히 기억나지 않음",
        "상속/증여/재개발 등으로 취득일 판단이 복잡함",
        "공동명의 또는 지분관계가 복잡함",
        "토지와 건물이 함께 있는 물건임",
        "주택과 상가가 함께 있는 겸용주택임",
    ])


def build_03(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 세대 기준 확인")
    add_field_table(doc, [
        ("현재 주민등록상 세대주", ""),
        ("배우자 여부", "□ 있음 □ 없음 □ 확인 필요"),
        ("배우자와 같은 세대 여부", "□ 같은 세대 □ 별도 세대 □ 확인 필요"),
        ("본인 명의 주택 수", ""),
        ("배우자 명의 주택 수", ""),
        ("세대원 명의 주택 수", ""),
    ])
    add_section_heading(doc, "2. 주택 수에 영향을 줄 수 있는 권리")
    add_check_items(doc, [
        "분양권 있음",
        "입주권 있음",
        "상속주택 있음",
        "공동상속주택 있음",
        "임대사업자 등록 주택 있음",
        "오피스텔 보유",
        "농어촌주택 보유",
        "장기임대주택 보유",
        "주택 수 포함 여부가 불명확한 자산 있음",
    ])


def build_04(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 실제 거주 여부")
    add_field_table(doc, [
        ("실제 거주 여부", "□ 예 □ 아니오 □ 일부 기간 □ 모름"),
        ("최초 전입일", ""),
        ("최종 전출일", ""),
        ("총 실제 거주기간", ""),
        ("현재 거주 여부", "□ 거주 중 □ 전출 □ 임대 중 □ 기타"),
        ("주민등록초본 제출 가능 여부", "□ 가능 □ 불가능 □ 확인 필요"),
    ])
    add_section_heading(doc, "2. 거주 증빙자료")
    add_check_items(doc, [
        "주민등록초본",
        "주민등록등본",
        "관리비 납부내역",
        "전기/수도/가스 사용내역",
        "임대차계약서",
        "기타 실제 거주 입증자료",
    ])


def build_05(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 취득가액 확인 상태")
    add_field_table(doc, [
        ("취득가액을 정확히 알고 있습니까?", "□ 예 □ 아니오 □ 대략 기억 □ 모름"),
        ("기억하는 취득가액", ""),
        ("매수계약서 보유 여부", "□ 있음 □ 없음 □ 확인 필요"),
        ("계약서 원본/사본 여부", "□ 원본 □ 사본 □ 없음 □ 확인 필요"),
        ("금융거래 내역 보유 여부", "□ 있음 □ 없음 □ 확인 필요"),
        ("취득세 납부자료 보유 여부", "□ 있음 □ 없음 □ 확인 필요"),
        ("취득세 자료에 과세표준/취득가액 표시 여부", "□ 표시됨 □ 표시 안 됨 □ 확인 필요"),
        ("등기부등본으로 취득일 확인 가능 여부", "□ 가능 □ 불가능 □ 확인 필요"),
    ])
    add_section_heading(doc, "2. 취득가액 불명확 체크")
    add_check_items(doc, [
        "취득가액을 정확히 기억하지 못함",
        "매수계약서가 없음",
        "금융거래 내역이 없음",
        "취득세 자료가 없음",
        "취득세 자료는 있으나 과세표준/취득가액 표시 여부를 모름",
        "오래전에 취득함",
        "공시지가/기준시가 확인 필요",
        "토지등급 자료 확인 필요",
        "매매사례가액 검토 필요",
        "감정가액 검토 필요",
        "환산취득가액 검토 필요",
        "의제취득일 검토 필요",
    ])
    add_section_heading(doc, "3. 필요경비 증빙 확인")
    add_field_table(doc, [
        ("취득세", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("중개수수료-취득", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("법무사비", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("인테리어/수리비", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("확장공사/설비공사", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("양도 중개수수료", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
        ("기타 양도비", "□ 있음 □ 없음 □ 확인 필요 / 금액:"),
    ])


def build_06(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 주택 수/비과세 관련")
    add_check_items(doc, [
        "2주택 이상 보유",
        "일시적 2주택 가능성",
        "배우자 명의 주택 있음",
        "세대원 명의 주택 있음",
        "분양권 있음",
        "입주권 있음",
        "상속주택 있음",
        "공동상속주택 있음",
        "조정대상지역 취득 여부 확인 필요",
        "실제 거주기간 불명확",
    ])
    add_section_heading(doc, "2. 취득 경위/거래상황 관련")
    add_check_items(doc, [
        "상속 취득",
        "증여 취득",
        "부담부증여 관련",
        "재개발/재건축",
        "조합원입주권 전환",
        "공동명의",
        "지분 일부 양도",
        "가족 간 거래",
        "특수관계자 거래",
        "이미 계약 완료",
        "잔금일 임박",
        "기존 세무 상담 이력 있음",
    ])
    add_section_heading(doc, "3. 상세 설명")
    doc.add_paragraph("\n\n\n")


def build_07(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 동의 항목")
    add_check_items(doc, [
        "개인정보 수집 및 이용에 동의합니다.",
        "부동산 거래정보 수집 및 이용에 동의합니다.",
        "가족/세대원 관련 정보 수집 및 이용에 동의합니다.",
        "세무사 검토가 필요한 경우 자료 전달에 동의합니다.",
        "AI 사전진단 처리에 동의합니다.",
        "녹음파일 제출은 선택사항임을 확인했습니다.",
        "본 리포트가 세무 판단 확정 또는 신고서가 아님을 확인했습니다.",
    ])
    add_signature_block(doc)


def build_08(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 기본 필수자료")
    add_check_items(doc, [
        "매도 부동산 등기부등본",
        "매수계약서",
        "매도계약서 또는 매도 예정금액 자료",
        "취득세 납부확인서",
        "주민등록초본",
        "주민등록등본",
    ])
    add_section_heading(doc, "2. 필요경비 자료")
    add_check_items(doc, [
        "취득 당시 중개수수료 영수증",
        "양도 당시 중개수수료 영수증",
        "법무사비 영수증",
        "인테리어/수리비 증빙",
        "공사계약서",
        "세금계산서/카드영수증/계좌이체 내역",
    ])
    add_section_heading(doc, "3. 예외사건 추가자료")
    add_check_items(doc, [
        "상속 관련 서류",
        "증여계약서",
        "분양계약서",
        "조합원입주권 관련 서류",
        "임대사업자 등록 관련 자료",
        "감정평가서",
        "기준시가/공시지가 확인자료",
        "토지등급 관련 자료",
    ])


def build_09(doc):
    add_common_header_table(doc)
    add_section_heading(doc, "1. 고객 최종 확인")
    add_check_items(doc, [
        "매도 부동산 정보가 맞습니다.",
        "취득일/양도예정일 정보가 맞습니다.",
        "취득가액 및 필요경비 자료 제출 상태가 맞습니다.",
        "본인/배우자/세대원 주택 보유 현황을 사실대로 작성했습니다.",
        "상속/증여/분양권/입주권 등 예외사항을 사실대로 표시했습니다.",
        "모르는 항목은 임의로 작성하지 않고 모름/확인 필요로 표시했습니다.",
        "본 리포트가 세무 판단 확정 또는 신고서가 아님을 확인했습니다.",
        "최종 판단은 세무사 검토가 필요할 수 있음을 확인했습니다.",
    ])
    add_signature_block(doc)


def build_10(doc):
    add_common_header_table(doc)
    doc.add_paragraph("1. 본인은 양도소득세 사전진단 및 신고에 필요한 서류와 정보를 성실히 제공하였음을 확인합니다.")
    doc.add_paragraph("2. 사실과 다르거나 누락된 정보로 인해 발생하는 추가 세금 및 가산세 등의 불이익은 본인에게 책임이 있음을 확인합니다.")
    add_signature_block(doc)


def build_11(doc):
    add_common_header_table(doc)
    doc.add_paragraph("1. 본인은 양도소득세 신고대행을 세무사에게 위임할 수 있음을 확인합니다.")
    doc.add_paragraph("2. 세무사 신고대행은 사전진단 리포트 또는 세무사 검토 패키지와 별도 업무임을 확인합니다.")
    doc.add_paragraph("3. 실제 신고대행 진행 시 세무사 사무실 기준에 따라 아래 자료가 추가로 요구될 수 있음을 확인합니다.")
    add_check_items(doc, [
        "홈택스 세무대리 수임동의",
        "양도소득세 신고대행 위임장",
        "신분증 사본",
        "인감증명서 또는 본인서명사실확인서",
        "인감도장 날인 또는 전자서명",
        "매매계약서",
        "취득계약서",
        "필요경비 증빙",
        "주민등록초본/등본",
        "세무사가 요청하는 기타 자료"
    ])
    doc.add_paragraph("4. 본인은 세무사가 본인확인 및 위임확인이 충분하지 않다고 판단하는 경우 신고대행을 보류하거나 거절할 수 있음을 확인합니다.")
    doc.add_paragraph("5. 본인은 세무대리 수임동의 또는 위임장 제출이 완료되지 않으면 실제 신고대행이 진행되지 않을 수 있음을 확인합니다.")
    doc.add_paragraph("6. 본인은 양도세 신고 내용의 최종 확인 책임이 본인에게 있음을 이해하고, 세무사가 작성한 신고내용을 제출 전 확인해야 함을 확인합니다.")
    add_signature_block(doc)


def build_12(doc):
    add_common_header_table(doc)
    doc.add_paragraph("고액 부동산과 복잡 사건에서 세무사 수수료 분쟁을 방지하기 위한 확인서입니다.")
    doc.add_paragraph("1. AI 양도세 사전진단 리포트 비용, 세무사 검토 패키지 비용, 실제 신고대행 수수료는 서로 다른 비용임을 확인합니다.")
    doc.add_paragraph("2. 기본 세무사 검토 패키지 비용 200,000원에는 실제 양도소득세 신고대행 수수료가 포함되지 않을 수 있습니다.")
    doc.add_paragraph("3. 실제 신고대행 수수료는 아래 요소에 따라 세무사가 별도 견적합니다.")
    add_check_items(doc, [
        "양도가액",
        "취득가액 불명확 여부",
        "필요경비 증빙 수량",
        "주택 수 판단 난이도",
        "조정대상지역 여부",
        "상속/증여/부담부증여 여부",
        "재개발·재건축/조합원입주권/분양권 여부",
        "1주택에서 2개 이상 입주권/신축주택 발생 여부",
        "양도 물건 수",
        "세무조사 위험 또는 쟁점 수"
    ])
    doc.add_paragraph("4. 고액 부동산 또는 복잡 사건은 추가 검토비가 발생할 수 있습니다.")
    doc.add_paragraph("5. 고객이 자료를 늦게 제출하거나 누락자료가 많을 경우 신고기한 임박 추가비 또는 긴급처리비가 발생할 수 있습니다.")
    doc.add_paragraph("6. 세무사는 자료 부족, 쟁점 과다, 고객 확인 미이행, 위임 미완료 시 신고대행을 거절할 수 있습니다.")
    doc.add_paragraph("7. 신고 완료 후 수정신고, 경정청구, 세무서 소명, 세무조사 대응은 별도 업무이며 별도 수수료가 발생할 수 있습니다.")
    add_section_heading(doc, "수수료 표")
    add_field_table(doc, [
        ("AI 사전진단 리포트", "100,000원 (포함/별도)"),
        ("세무사 검토 패키지", "200,000원 (포함/별도)"),
        ("양도세 신고대행", "세무사 견적 (별도)"),
        ("복잡사건 추가 검토", "세무사 견적 (별도)"),
        ("수정신고/경정청구/소명대응", "세무사 견적 (별도)"),
    ])
    add_signature_block(doc)


def build_13(doc):
    add_common_header_table(doc)
    doc.add_paragraph("세무사가 실제 신고하기 전에 고객이 최종 신고내용을 확인했음을 남기는 확인서입니다.")
    doc.add_paragraph("1. 본인은 세무사가 작성한 양도소득세 신고내용을 제출 전 확인했습니다.")
    doc.add_paragraph("2. 본인은 아래 항목을 확인했습니다.")
    add_check_items(doc, [
        "양도 부동산 주소",
        "취득일",
        "양도일",
        "취득가액",
        "양도가액",
        "필요경비",
        "보유기간",
        "거주기간",
        "주택 수",
        "분양권/입주권/상속주택 여부",
        "조정대상지역 여부",
        "세액 계산 결과",
        "납부세액",
        "납부기한"
    ])
    doc.add_paragraph("3. 본인은 신고내용 중 모르는 부분이나 이견이 있는 부분을 세무사에게 즉시 알려야 함을 확인합니다.")
    doc.add_paragraph("4. 본인이 최종 확인한 후 제출된 신고내용이라도, 고객이 제공한 자료가 사실과 다르거나 누락된 경우 추가세액, 가산세, 납부지연가산세 등이 발생할 수 있음을 확인합니다.")
    doc.add_paragraph("5. 본인은 신고 후 세무서에서 추가자료 제출 또는 소명을 요청할 수 있음을 확인합니다.")
    add_signature_block(doc)


BUILDERS = {
    "01_client_tax_consult_application.docx": build_01,
    "02_property_basic_info.docx": build_02,
    "03_household_house_status.docx": build_03,
    "04_residence_period_check.docx": build_04,
    "05_acquisition_price_expense_check.docx": build_05,
    "06_exception_issue_checklist.docx": build_06,
    "07_privacy_and_tax_accountant_consent.docx": build_07,
    "08_required_documents_checklist.docx": build_08,
    "09_final_client_confirmation.docx": build_09,
    "10_client_truthful_disclosure_notice.docx": build_10,
    "11_tax_filing_delegation_and_identity_check.docx": build_11,
    "12_tax_accountant_fee_scope_agreement.docx": build_12,
    "13_final_tax_filing_confirmation.docx": build_13,
}


def create_docx(filename, title, out_dir):
    doc = Document()
    set_doc_style(doc)
    add_title(doc, title)
    add_notice(doc)

    builder = BUILDERS.get(filename)
    if builder:
        builder(doc)

    add_notice(doc)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / filename
    doc.save(path)
    return path


def create_all_docx(out_dir):
    paths = []
    for filename, title in FORM_FILES.items():
        paths.append(create_docx(filename, title, out_dir))
    return paths


def copy_to_case(case_id, template_paths):
    case_dir = CASES_DIR / case_id
    if not case_dir.exists():
        raise FileNotFoundError(f"사건 폴더를 찾을 수 없습니다: {case_dir}")

    target_dir = case_dir / "09_print_forms"
    target_dir.mkdir(exist_ok=True)

    copied = []
    for src in template_paths:
        dst = target_dir / src.name
        shutil.copy2(src, dst)
        copied.append(dst)

    return copied


def main():
    parser = argparse.ArgumentParser(description="TaxCaseManager 고객 작성용 DOCX 출력 양식 생성기")
    parser.add_argument("--case-id", help="사건번호를 입력하면 사건 폴더에도 DOCX 양식을 복사")
    args = parser.parse_args()

    try:
        template_paths = create_all_docx(TEMPLATE_DOCX_DIR)
        output_paths = create_all_docx(OUTPUT_DOCX_DIR)

        print("[SUCCESS] DOCX 출력 양식 생성 완료")
        print(f"[TEMPLATE_DIR] {TEMPLATE_DOCX_DIR}")
        print(f"[OUTPUT_DIR] {OUTPUT_DOCX_DIR}")
        print(f"[COUNT] {len(output_paths)}")

        copied = []
        if args.case_id:
            copied = copy_to_case(args.case_id, template_paths)
            print("")
            print("[SUCCESS] 사건 폴더 DOCX 복사 완료")
            print(f"[CASE_ID] {args.case_id}")
            print(f"[CASE_PRINT_DIR] {CASES_DIR / args.case_id / '09_print_forms'}")
            print(f"[COUNT] {len(copied)}")

        write_log(f"created_docx={len(output_paths)} case_id={args.case_id or ''} copied={len(copied)}")

    except Exception as e:
        write_log(f"ERROR {type(e).__name__}: {e}")
        print("[FAIL] DOCX 출력 양식 생성 실패")
        print(f"[ERROR] {type(e).__name__}: {e}")


if __name__ == "__main__":
    main()
