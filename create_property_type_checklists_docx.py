import argparse
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(r"C:\TaxCaseManager")
TEMPLATE_DOCX_DIR = BASE_DIR / "templates" / "property_type_checklists_docx"
OUTPUT_DOCX_DIR = BASE_DIR / "outputs" / "property_type_checklists_docx"
CASES_DIR = BASE_DIR / "cases"
LOGS_DIR = BASE_DIR / "logs"

NOTICE_TEXT = "본 서류는 양도세 사전진단을 위한 기초자료이며, 최종 세무 판단은 세무사 검토가 필요합니다."
STRONG_NOTICE_TEXT = "재개발·재건축·조합원입주권·분양권 사건은 세무 쟁점이 복잡하므로 부동산 사무실에서 세무 판단을 확정하지 않습니다. 본 사건은 세무사 검토가 강력히 권장되거나 필수입니다."

FORM_FILES = {
    "simple": ("01_simple_apartment_house_checklist.docx", "일반 아파트/단독/다세대 간편 체크리스트"),
    "dagagu": ("02_dagagu_multifamily_checklist.docx", "다가구/다세대/임대 주택 체크리스트"),
    "mixed": ("03_mixed_use_house_checklist.docx", "겸용주택/상가주택 체크리스트"),
    "land": ("04_land_farmland_forest_commercial_checklist.docx", "토지/농지/임야/상가 체크리스트"),
    "redevelopment": ("05_redevelopment_reconstruction_checklist.docx", "재개발/재건축/입주권/분양권 체크리스트"),
    "inheritance": ("06_inheritance_gift_checklist.docx", "상속/증여/부담부증여 체크리스트"),
    "tax_required": ("07_tax_accountant_required_checklist.docx", "세무사 필수 검토 체크리스트"),
}

def write_log(message):
    LOGS_DIR.mkdir(exist_ok=True)
    log_path = LOGS_DIR / "property_type_checklists_docx_log.txt"
    now = datetime.now().isoformat(timespec="seconds")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{now}] {message}\n")

def set_cell_shading(cell, fill="EDEDED"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_doc_style(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)

def add_title(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(16)
    doc.add_paragraph("")

def add_notice(doc, strong=False):
    p = doc.add_paragraph()
    text = STRONG_NOTICE_TEXT if strong else NOTICE_TEXT
    run = p.add_run("※ " + text)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(9)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(12)

def add_field_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "항목", bold=True)
    set_cell_text(hdr[1], "작성 내용", bold=True)
    set_cell_shading(hdr[0])
    set_cell_shading(hdr[1])

    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], value)

    doc.add_paragraph("")
    return table

def add_check_items(doc, items):
    for item in items:
        p = doc.add_paragraph()
        run = p.add_run(f"□ {item}")
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        run.font.size = Pt(10.5)
    doc.add_paragraph("")

def add_signature_block(doc):
    doc.add_paragraph("")
    add_field_table(doc, [
        ("고객명", "____________________________"),
        ("서명", "____________________________"),
        ("작성일", "________년 ________월 ________일"),
    ])

def add_common_header(doc):
    add_field_table(doc, [
        ("고객명", ""),
        ("연락처", ""),
        ("부동산 주소", ""),
    ])

def build_simple(doc):
    p = doc.add_paragraph()
    run = p.add_run("※ 주의: 본 체크리스트는 일반 주택 사전진단용입니다. 2주택, 상속, 증여, 재개발/재건축, 취득가액 불명확 사건은 추가 체크리스트가 필요합니다.")
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(9)
    doc.add_paragraph("")
    
    add_common_header(doc)
    add_section_heading(doc, "1. 기본 정보")
    add_field_table(doc, [
        ("부동산 종류", "□ 일반 아파트 □ 일반 단독주택 □ 일반 다세대/빌라"),
        ("취득일", ""),
        ("취득가액", ""),
        ("양도 예정일", ""),
        ("양도 예정가액", ""),
        ("실제 거주기간", ""),
    ])
    add_section_heading(doc, "2. 주택 수 및 증빙 자료")
    add_check_items(doc, [
        "본인 주택 수: ____채",
        "배우자 주택 있음",
        "세대원 주택 있음",
        "분양권/입주권 있음",
        "상속주택 있음",
        "매수계약서 보유",
        "취득세 자료 보유",
        "중개수수료 자료 보유",
        "세무사 검토 희망",
    ])

def build_dagagu(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 다가구/다세대/임대 현황")
    add_field_table(doc, [
        ("다가구/다세대 확인", "□ 다가구(단독주택) □ 다세대(공동주택) □ 확인 필요"),
        ("건축물대장상 용도", ""),
        ("실제 사용 용도", ""),
        ("총 호실 수", ""),
        ("주인 거주 여부", "□ 거주 중 □ 비거주"),
    ])
    add_section_heading(doc, "2. 상세 체크")
    add_check_items(doc, [
        "각 호실 임대 중",
        "층별 사용 현황 확인 필요",
        "임대차계약서 보유",
        "주택 수 판단 쟁점 있음",
    ])

def build_mixed(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 겸용/상가주택 현황")
    add_field_table(doc, [
        ("주택 면적", ""),
        ("상가 면적", ""),
        ("건축물대장상 용도", ""),
        ("주택 면적이 상가 면적보다 큰지", "□ 예 □ 아니오 □ 모름"),
    ])
    add_section_heading(doc, "2. 상세 체크")
    add_check_items(doc, [
        "주택 부분 실제 사용 중",
        "상가 부분 임대 중",
        "층별 용도 확인 필요",
        "토지/건물 가액 구분 가능",
        "세무사 검토 필요",
    ])

def build_land(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 토지/상가 등 기본정보")
    add_field_table(doc, [
        ("자산 종류", "□ 토지 □ 농지 □ 임야 □ 상가"),
        ("지목", ""),
        ("실제 사용 현황", ""),
        ("취득일", ""),
        ("취득가액", ""),
        ("양도 예정가액", ""),
    ])
    add_section_heading(doc, "2. 상세 체크")
    add_check_items(doc, [
        "공시지가/기준시가 자료 보유",
        "비사업용 토지 가능성 있음",
        "농지 자경 (직접 농사 지음)",
        "임야 보유 목적: _____________",
        "상가 임대 중",
        "필요경비 자료 보유",
    ])

def build_redevelopment(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 재개발/재건축/입주권/분양권 정보")
    add_field_table(doc, [
        ("기존 주택 최초 취득일", ""),
        ("사업시행인가일", ""),
        ("관리처분계획인가일", ""),
        ("관리처분인가일 당시 조정대상지역 여부", "□ 예 □ 아니오 □ 모름"),
        ("관리처분인가일 당시 보유 주택 수", ""),
        ("기존 주택 멸실일", ""),
        ("입주권/분양권 수", ""),
        ("1개 주택이 2개 입주권/분양권으로 전환되었는지 여부", "□ 예 □ 아니오"),
        ("신축 아파트 수", ""),
        ("각 물건별 양도 예정가액", ""),
        ("각 물건별 양도 예정일", ""),
    ])
    add_section_heading(doc, "2. 추가 체크")
    add_check_items(doc, [
        "어느 물건을 먼저 양도할지: _____________",
        "실거주 요건 검토 필요 여부",
        "대체주택 취득/양도함",
        "세무사 필수 검토",
    ])

def build_inheritance(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 상속/증여 현황")
    add_field_table(doc, [
        ("상속 취득 여부", "□ 예 □ 아니오"),
        ("상속개시일", ""),
        ("피상속인 취득일", ""),
        ("상속 지분", ""),
        ("공동상속 여부", "□ 예 □ 아니오"),
        ("증여 취득 여부", "□ 예 □ 아니오"),
        ("증여일", ""),
        ("증여가액", ""),
    ])
    add_section_heading(doc, "2. 상세 체크")
    add_check_items(doc, [
        "부담부증여",
        "채무 인수",
        "상속/증여 당시 평가자료 보유",
        "세무사 검토 필요",
    ])

def build_tax_required(doc):
    add_common_header(doc)
    add_section_heading(doc, "1. 세무사 필수 검토 사유")
    add_check_items(doc, [
        "취득가액 불명확",
        "매수계약서 없음",
        "금융자료 없음",
        "2주택 이상",
        "배우자/세대원 주택 미확인",
        "상속/증여",
        "재개발/재건축",
        "입주권/분양권",
        "토지/농지/임야",
        "비사업용 토지 가능성",
        "고가주택",
        "잔금일 임박",
        "신고기한 임박",
        "기존 세무상담 내용과 불일치",
    ])

BUILDERS = {
    "simple": build_simple,
    "dagagu": build_dagagu,
    "mixed": build_mixed,
    "land": build_land,
    "redevelopment": build_redevelopment,
    "inheritance": build_inheritance,
    "tax_required": build_tax_required,
}

def create_docx(type_key, out_dir):
    filename, title = FORM_FILES[type_key]
    doc = Document()
    set_doc_style(doc)
    add_title(doc, title)
    
    strong_notice = (type_key == "redevelopment")
    add_notice(doc, strong=strong_notice)

    builder = BUILDERS.get(type_key)
    if builder:
        builder(doc)

    add_signature_block(doc)
    add_notice(doc, strong=False)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / filename
    doc.save(path)
    return path

def copy_to_case(case_id, template_paths):
    case_dir = CASES_DIR / case_id
    if not case_dir.exists():
        raise FileNotFoundError(f"사건 폴더를 찾을 수 없습니다: {case_dir}")

    target_dir = case_dir / "09_print_forms" / "property_type_checklists"
    target_dir.mkdir(parents=True, exist_ok=True)

    copied = []
    for src in template_paths:
        dst = target_dir / src.name
        shutil.copy2(src, dst)
        copied.append(dst)

    return copied

def main():
    parser = argparse.ArgumentParser(description="TaxCaseManager 물건 유형별 간편 체크리스트 DOCX 생성기")
    parser.add_argument("--all", action="store_true", help="모든 타입 생성")
    parser.add_argument("--type", choices=FORM_FILES.keys(), help="특정 타입 생성")
    parser.add_argument("--case-id", help="사건번호를 입력하면 사건 폴더에도 DOCX 복사")
    args = parser.parse_args()

    types_to_build = []
    if args.all:
        types_to_build = list(FORM_FILES.keys())
    elif args.type:
        types_to_build = [args.type]
    else:
        parser.print_help()
        return

    try:
        template_paths = []
        output_paths = []
        for t in types_to_build:
            t_path = create_docx(t, TEMPLATE_DOCX_DIR)
            o_path = create_docx(t, OUTPUT_DOCX_DIR)
            template_paths.append(t_path)
            output_paths.append(o_path)

        print("[SUCCESS] DOCX 체크리스트 생성 완료")
        print(f"[TEMPLATE_DIR] {TEMPLATE_DOCX_DIR}")
        print(f"[OUTPUT_DIR] {OUTPUT_DOCX_DIR}")
        print(f"[COUNT] {len(output_paths)}")

        copied = []
        if args.case_id:
            copied = copy_to_case(args.case_id, template_paths)
            print("")
            print("[SUCCESS] 사건 폴더 DOCX 복사 완료")
            print(f"[CASE_ID] {args.case_id}")
            print(f"[CASE_PRINT_DIR] {CASES_DIR / args.case_id / '09_print_forms' / 'property_type_checklists'}")
            print(f"[COUNT] {len(copied)}")

        write_log(f"types={','.join(types_to_build)} count={len(output_paths)} case_id={args.case_id or ''} copied={len(copied)}")

    except Exception as e:
        write_log(f"ERROR {type(e).__name__}: {e}")
        print("[FAIL] DOCX 생성 실패")
        print(f"[ERROR] {type(e).__name__}: {e}")

if __name__ == "__main__":
    main()
