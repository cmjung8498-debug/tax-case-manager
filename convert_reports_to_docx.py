import argparse
import sys
from datetime import datetime
from pathlib import Path

try:
    import docx
except ImportError:
    print("[ERROR] python-docx 패키지가 필요합니다. pip install python-docx")
    sys.exit(1)

BASE_DIR = Path(r"C:\TaxCaseManager")
CASES_DIR = BASE_DIR / "cases"
LOGS_DIR = BASE_DIR / "logs"


def write_log(message: str):
    LOGS_DIR.mkdir(exist_ok=True)
    log_path = LOGS_DIR / "report_docx_log.txt"
    now = datetime.now().isoformat(timespec="seconds")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{now}] {message}\n")


def convert_md_to_docx(md_path: Path, docx_path: Path):
    doc = docx.Document()
    
    # 기본 스타일 조정 (선택 사항)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Malgun Gothic'
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_table = False
    
    for line in lines:
        line = line.rstrip('\n')
        
        # 빈 줄은 건너뛰거나 문단 추가
        if not line.strip():
            doc.add_paragraph("")
            continue
            
        # 헤딩 처리
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            
        # 리스트 아이템 처리
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style='List Bullet')
            
        # 테이블 처리 (매우 단순한 구현)
        elif line.startswith("|") and line.endswith("|"):
            doc.add_paragraph(line) # 파이썬 docx로 깔끔한 표 처리는 복잡하므로 단순 텍스트로 일단 처리
        else:
            doc.add_paragraph(line)
            
    doc.save(docx_path)


def main():
    parser = argparse.ArgumentParser(description="Markdown 리포트를 DOCX로 변환")
    parser.add_argument("case_id", help="사건번호")
    args = parser.parse_args()

    case_id = args.case_id
    case_dir = CASES_DIR / case_id
    report_dir = case_dir / "06_reports"
    docx_dir = case_dir / "06_reports_docx"

    try:
        if not report_dir.exists():
            print(f"[FAIL] 06_reports 폴더가 없습니다: {report_dir}")
            write_log(f"FAIL case_id={case_id} msg='06_reports not found'")
            sys.exit(1)

        docx_dir.mkdir(parents=True, exist_ok=True)
        count = 0

        for md_file in report_dir.glob("*.md"):
            docx_file = docx_dir / (md_file.stem + ".docx")
            convert_md_to_docx(md_file, docx_file)
            count += 1

        print("[SUCCESS] DOCX 리포트 변환 완료")
        print(f"[CASE_ID] {case_id}")
        print(f"[DOCX_DIR] {docx_dir}")
        print(f"[DOCX_COUNT] {count}")

        write_log(f"SUCCESS case_id={case_id} converted={count} files")

    except Exception as e:
        print("[FAIL] DOCX 변환 실패")
        print(f"[ERROR] {type(e).__name__}: {e}")
        write_log(f"ERROR case_id={case_id} msg='{type(e).__name__}: {e}'")
        sys.exit(1)


if __name__ == "__main__":
    main()
